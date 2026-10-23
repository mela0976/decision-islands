#!/usr/bin/env python3
"""Build the static article catalog from a JSON edition manifest.

Runtime dependencies: openpyxl (XLSX), pypdf (PDF), Pillow (WebP conversion).
All source paths in the manifest are relative to the project root. Output paths
are relative to the website root. The importer validates the whole edition
before writing and replaces only that edition's generated files/catalog entry.
"""
from __future__ import annotations

import argparse
import html
import json
import math
import re
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from PIL import Image, ImageOps
from pypdf import PdfReader, PdfWriter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from reading_order import reconcile_reading_order

CHAPTERS = {"intro", "people", "model", "ai"}


class ImportErrorDetail(ValueError):
    pass


def _fail(message: str) -> None:
    raise ImportErrorDetail(message)


def _read_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        _fail(f"cannot read JSON {path}: {exc}")
    if not isinstance(value, dict):
        _fail(f"manifest must be a JSON object: {path}")
    return value


def _source(root: Path, rel: str, label: str) -> Path:
    if not isinstance(rel, str) or not rel or Path(rel).is_absolute():
        _fail(f"{label} must be a project-relative path")
    path = (root / rel).resolve()
    if not path.is_relative_to(root.resolve()):
        _fail(f"{label} escapes project root: {rel}")
    if not path.is_file():
        _fail(f"missing {label}: {rel}")
    return path


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _workbook_metadata(root: Path, manifest: dict[str, Any]) -> dict[int, dict[str, Any]]:
    path = _source(root, manifest["metadataWorkbook"], "metadata workbook")
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        _fail(f"cannot open metadata workbook {path}: {exc}")
    sheet_name = manifest.get("metadataSheet")
    if sheet_name not in wb.sheetnames:
        _fail(f"metadata sheet not found: {sheet_name}")
    ws = wb[sheet_name]
    header_row = manifest.get("headerRow", 1)
    if not isinstance(header_row, int) or header_row < 1:
        _fail("headerRow must be a positive integer")
    headers = {_cell_text(cell.value): cell.column for cell in ws[header_row] if _cell_text(cell.value)}
    columns = manifest.get("metadataColumns")
    required = {key: columns.get(key) for key in ("title", "summary", "authors", "highlights")} if isinstance(columns, dict) else {}
    if set(required) != {"title", "summary", "authors", "highlights"} or any(v not in headers for v in required.values()):
        _fail("metadataColumns must name existing title, summary, authors and highlights headers")
    out: dict[int, dict[str, Any]] = {}
    for spec in manifest.get("articles", []):
        row_num = spec.get("metadataRow")
        if not isinstance(row_num, int) or row_num <= header_row:
            _fail(f"invalid metadataRow for {spec.get('id')}")
        values = {key: _cell_text(ws.cell(row_num, headers[col]).value) for key, col in required.items()}
        if not values["title"] or not values["authors"] or not values["summary"]:
            _fail(f"incomplete workbook metadata at row {row_num}")
        values["highlights"] = _parse_highlights(values["highlights"])
        if len(values["highlights"]) != 3:
            _fail(f"expected three workbook highlights at row {row_num}, found {len(values['highlights'])}")
        out[row_num] = values
    wb.close()
    return out


def _parse_highlights(raw: str) -> list[str]:
    items: list[str] = []
    for line in raw.replace("\r", "\n").split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("*") or line.startswith("•") or line.startswith("●"):
            items.append(line.lstrip("*•● ").strip())
        elif items:
            items[-1] += " " + line
    return [x for x in items if x]


def _tag_text(element: ET.Element) -> str:
    return element.tag.rsplit("}", 1)[-1]


def _extract_idml(path: Path, prefixes: list[str]) -> list[str]:
    try:
        with zipfile.ZipFile(path) as archive:
            design = ET.fromstring(archive.read("designmap.xml"))
            story_ids = design.attrib.get("StoryList", "").split()
            stories = {Path(name).stem.removeprefix("Story_"): name for name in archive.namelist() if name.startswith("Stories/") and name.endswith(".xml")}
            paragraphs: list[str] = []
            for story_id in story_ids:
                name = stories.get(story_id)
                if not name:
                    continue
                root = ET.fromstring(archive.read(name))
                for p in (e for e in root.iter() if _tag_text(e) == "ParagraphStyleRange"):
                    style = p.attrib.get("AppliedParagraphStyle", "").split("/")[-1]
                    if not any(style.startswith(prefix) for prefix in prefixes):
                        continue
                    pieces: list[str] = []
                    for el in p.iter():
                        tag = _tag_text(el)
                        if tag == "Content" and el.text:
                            pieces.append(el.text)
                        elif tag == "Br":
                            pieces.append("\n")
                        elif tag == "ForcedLineBreak":
                            pieces.append(" ")
                    for run in "".join(pieces).split("\n"):
                        text = re.sub(r"\s+", " ", run).strip()
                        text = text.replace("\ufffc", "").strip()
                        if text and (not paragraphs or paragraphs[-1] != text):
                            paragraphs.append(text)
            if not paragraphs:
                _fail(f"no readable paragraphs in IDML: {path}")
            return paragraphs
    except (OSError, zipfile.BadZipFile, KeyError, ET.ParseError) as exc:
        _fail(f"cannot parse IDML {path}: {exc}")


def _article_html(article: dict[str, Any], edition_title: str, body_source: str) -> str:
    e = html.escape
    body = "\n".join(f"<p>{e(p)}</p>" for p in article["paragraphs"])
    highlights = "\n".join(f"<li>{e(x)}</li>" for x in article["highlights"])
    supplementary = "".join(f"<p>{e(p)}</p>" for p in article.get("supplementaryParagraphs", []))
    supplementary_html = f'<details><summary>圖表、註解與補充文字</summary><p class="meta">以下文字見於原刊的多個位置，集中保留供查閱；圖表關係請對照 PDF。</p>{supplementary}</details>' if supplementary else ''
    pages = f"印刷頁 {article['printedStart']}–{article['printedEnd']}；完整 PDF 第 {article['sourcePages'][0]}–{article['sourcePages'][-1]} 頁"
    chapter_labels = {"intro": "啟航", "people": "人心之島", "model": "洞察之島", "ai": "行動之島"}
    return f'''<!doctype html>
<html lang="zh-Hant"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><meta name="theme-color" content="#faf8f2">
<title>{e(article['title'])}｜{e(edition_title)}</title><meta name="description" content="{e(article['summary'])}">
<style>
:root{{color-scheme:light;--paper:#faf8f2;--ink:#263c3e;--muted:#63706b;--accent:#9a5737;--rule:#ddd8ca}}
*{{box-sizing:border-box}}body{{max-width:720px;margin:0 auto;padding:1.5rem 1.4rem 4rem;background:var(--paper);color:var(--ink);font:1.12rem/2 "Noto Serif TC","Source Han Serif TC","Songti TC","PMingLiU",serif}}
nav{{display:flex;gap:1.1rem;flex-wrap:wrap;padding:.4rem 0 1.2rem;border-bottom:1px solid var(--rule);font: .9rem/1.6 system-ui,sans-serif}}a{{color:var(--accent);text-underline-offset:.2em}}
header{{padding:2.2rem 0 1.6rem;border-bottom:1px solid var(--rule)}}h1{{max-width:16ch;margin:.35rem 0 1rem;font-size:clamp(2rem,7vw,3.25rem);line-height:1.32;letter-spacing:.015em}}
.chapter{{color:var(--accent);font: .82rem/1.5 system-ui,sans-serif;letter-spacing:.13em}}.meta,footer{{color:var(--muted);font: .88rem/1.7 system-ui,sans-serif}}.summary{{margin-top:1.35rem;font-size:1.08rem;line-height:1.85}}
h2{{margin:2.4rem 0 .6rem;font-size:1.35rem}}section{{margin-top:1.8rem}}li{{margin:.45rem 0}}.body p{{margin:1.25em 0;text-wrap:pretty}}footer{{margin-top:3rem;padding-top:1.25rem;border-top:1px solid var(--rule)}}
@media(max-width:520px){{body{{padding:1rem 1rem 3rem;font-size:1.05rem;line-height:1.9}}header{{padding-top:1.7rem}}h1{{font-size:2rem}}}}
</style></head>
<body><nav aria-label="文章導覽"><a href="../../index.html">回到互動閱讀首頁</a><a href="index.html">回到 {article['year']} 文章目錄</a></nav><main><article><header><p class="chapter">{e(chapter_labels.get(article['chapter'], article['chapter']))} · {e(edition_title)}</p><h1>{e(article['title'])}</h1><p class="meta">{e(article['authors'])} · {e(pages)}</p><p class="summary">{e(article['summary'])}</p></header>
<section aria-labelledby="highlights"><h2 id="highlights">必看三要素</h2><ul>{highlights}</ul></section><section class="body" aria-label="文章文字">{body}</section>{supplementary_html}
<footer><p>編輯來源：標題、作者、摘要與重點依上稿資料；正文來源：{e(body_source)}。</p><p><a href="../../{e(article['sourcePdf'])}">下載本篇原刊 PDF</a>（{e(pages)}，含原刊跨頁註腳與參考文字）</p><p><a href="https://www.dentsu.com/uk/en/policies/responsible-disclosure-policy">電通負責任揭露政策</a></p></footer></article></main></body></html>'''


def _index_html(edition: dict[str, Any]) -> str:
    e = html.escape
    items = "\n".join(f'<li><a href="{e(a["id"])}.html">{e(a["title"])}</a> <span>— {e(a["authors"])}</span></li>' for a in edition["articles"])
    return f'''<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><meta name="theme-color" content="#faf8f2"><title>{e(edition['title'])}｜文章目錄</title>
<style>:root{{--paper:#faf8f2;--ink:#263c3e;--muted:#63706b;--accent:#9a5737;--rule:#ddd8ca}}*{{box-sizing:border-box}}body{{max-width:760px;margin:0 auto;padding:1.5rem 1.4rem 4rem;background:var(--paper);color:var(--ink);font:1.05rem/1.8 system-ui,"Noto Sans TC",sans-serif}}nav{{display:flex;gap:1rem;flex-wrap:wrap;padding:.4rem 0 1.2rem;border-bottom:1px solid var(--rule)}}a{{color:var(--accent);text-underline-offset:.2em}}h1{{max-width:20ch;margin:2rem 0 .5rem;font:clamp(2rem,7vw,3.25rem)/1.3 "Noto Serif TC","Songti TC","PMingLiU",serif}}.intro,.authors{{color:var(--muted)}}ol{{list-style:none;padding:0;margin:2rem 0}}li{{padding:1rem 0;border-top:1px solid var(--rule)}}li a{{display:block;font:1.25rem/1.5 "Noto Serif TC","Songti TC","PMingLiU",serif}}.authors{{font-size:.9rem}}footer{{margin-top:3rem;padding-top:1rem;border-top:1px solid var(--rule);font-size:.85rem}}@media(max-width:520px){{body{{padding:1rem}}}}</style></head>
<body><nav aria-label="文章導覽"><a href="../../index.html">回到互動閱讀首頁</a><a href="https://www.dentsu.com/uk/en/policies/responsible-disclosure-policy">電通負責任揭露政策</a></nav><main><h1>{e(edition['title'])}</h1><p class="intro">共 {len(edition['articles'])} 篇文章。此目錄與文章頁不需 JavaScript 即可閱讀。</p><ol>{items}</ol></main><footer>文章內容依出版資料編排；每篇均附原刊 PDF 來源頁碼與下載連結。</footer></body></html>'''


def _pdf_paragraphs(reader: PdfReader, pages: list[int]) -> list[str]:
    paragraphs: list[str] = []
    for page_no in pages:
        text = reader.pages[page_no - 1].extract_text() or ""
        paragraphs.extend(re.sub(r"\s+", " ", line).strip() for line in text.splitlines() if line.strip())
    if not paragraphs:
        _fail(f"PDF pages contain no extractable text: {pages}")
    return paragraphs


def _inline_metadata(spec: dict[str, Any]) -> dict[str, Any]:
    data = spec.get("metadata", spec)
    required = ("title", "authors", "summary", "highlights")
    if not isinstance(data, dict) or any(key not in data for key in required):
        _fail(f"PDF-only article {spec.get('id')} needs inline title, authors, summary and highlights")
    highlights = data["highlights"]
    if not isinstance(highlights, list) or not all(isinstance(x, str) and x.strip() for x in highlights):
        _fail(f"inline highlights must be a list of non-empty strings for {spec.get('id')}")
    result = {key: data[key] for key in required}
    if not all(isinstance(result[key], str) and result[key].strip() for key in ("title", "authors", "summary")):
        _fail(f"inline metadata contains an empty value for {spec.get('id')}")
    result["highlights"] = [x.strip() for x in highlights]
    return result


def _source_pages(spec: dict[str, Any], manifest: dict[str, Any]) -> list[int]:
    supplied = spec.get("sourcePages")
    if supplied is not None:
        if not isinstance(supplied, list) or not supplied or not all(isinstance(x, int) and x > 0 for x in supplied):
            _fail(f"sourcePages must be a non-empty list of positive 1-based pages for {spec.get('id')}")
        if supplied != sorted(set(supplied)):
            _fail(f"sourcePages must be unique and ascending for {spec.get('id')}")
        return supplied
    mapping = manifest.get("printPagePdfMapping")
    if not isinstance(mapping, dict) or not isinstance(mapping.get("divisor"), int) or mapping["divisor"] < 1 or not isinstance(mapping.get("offset"), int):
        _fail(f"{spec.get('id')} needs sourcePages or a manifest printPagePdfMapping")
    divisor, offset = mapping["divisor"], mapping["offset"]
    return list(range(spec["printedStart"] // divisor + offset, spec["printedEnd"] // divisor + offset + 1))


def _validate_existing_catalog(website_root: Path, year: int, article_ids: set[str]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    path = website_root / "public/data/catalog.json"
    if not path.exists():
        return {"schemaVersion": 1, "years": []}, []
    catalog = _read_json(path)
    if catalog.get("schemaVersion") != 1 or not isinstance(catalog.get("years"), list):
        _fail("existing catalog has an unsupported shape")
    others = [entry for entry in catalog["years"] if entry.get("year") != year]
    for entry in others:
        data_rel = entry.get("dataUrl")
        if not isinstance(data_rel, str):
            _fail(f"existing catalog entry lacks dataUrl: {entry}")
        data_path = website_root / "public" / data_rel
        if data_path.exists():
            data = _read_json(data_path)
            ids = {a.get("id") for a in data.get("articles", []) if isinstance(a, dict)}
            collision = ids & article_ids
            if collision:
                _fail(f"article IDs collide with year {entry.get('year')}: {sorted(collision)}")
    return catalog, others


def _validate_manifest(manifest: dict[str, Any], root: Path, website_root: Path) -> tuple[list[dict[str, Any]], dict[int, dict[str, Any]], PdfReader]:
    year = manifest.get("year")
    articles = manifest.get("articles")
    if not isinstance(year, int) or year < 1900 or year > 9999:
        _fail("year must be a four-digit integer")
    if not isinstance(manifest.get("title"), str) or not manifest["title"].strip():
        _fail("edition title is required")
    if not isinstance(articles, list) or not articles:
        _fail("manifest must include articles")
    seen: set[str] = set()
    previous_order = 0
    normalized = []
    for spec in articles:
        if not isinstance(spec, dict):
            _fail("each article manifest entry must be an object")
        ident = spec.get("id")
        if not isinstance(ident, str) or not re.fullmatch(rf"{year}-[A-Za-z0-9][A-Za-z0-9_-]*", ident) or ident in seen:
            _fail(f"duplicate or invalid article ID: {ident!r}")
        seen.add(ident)
        if not ident.startswith(f"{year}-"):
            _fail(f"article ID must be frozen with the edition year prefix: {ident}")
        order = spec.get("order")
        if not isinstance(order, int) or order <= previous_order:
            _fail(f"article order must be strictly increasing at {ident}")
        previous_order = order
        if spec.get("chapter") not in CHAPTERS:
            _fail(f"invalid chapter for {ident}: {spec.get('chapter')}")
        start, end = spec.get("printedStart"), spec.get("printedEnd")
        if not isinstance(start, int) or not isinstance(end, int) or start < 1 or end < start:
            _fail(f"invalid printed page range for {ident}: {start}–{end}")
        if spec.get("idmlSource"):
            _source(root, spec["idmlSource"], f"IDML source for {ident}")
        elif "paragraphs" in spec and (not isinstance(spec["paragraphs"], list) or not all(isinstance(x, str) for x in spec["paragraphs"])):
            _fail(f"paragraphs must be a string list for {ident}")
        _source_pages(spec, manifest)
        normalized.append(spec)
    has_workbook = bool(manifest.get("metadataWorkbook"))
    workbook: dict[int, dict[str, Any]] = {}
    if has_workbook:
        if len({a.get("metadataRow") for a in articles}) != len(articles):
            _fail("metadataRow values must be unique")
        workbook = _workbook_metadata(root, manifest)
        for spec in normalized:
            if spec.get("metadataRow") not in workbook:
                _fail(f"metadata missing for {spec['id']}")
    else:
        for spec in normalized:
            _inline_metadata(spec)
    pdf_path = _source(root, manifest.get("fullBookPdf", ""), "full-book PDF")
    try:
        pdf_reader = PdfReader(str(pdf_path))
    except Exception as exc:
        _fail(f"cannot read full-book PDF: {exc}")
    if len(pdf_reader.pages) < 1:
        _fail("source PDF is empty")
    pdf_last_allowed = manifest.get("maxSourcePage", len(pdf_reader.pages))
    if not isinstance(pdf_last_allowed, int) or pdf_last_allowed < 1 or pdf_last_allowed > len(pdf_reader.pages):
        _fail(f"maxSourcePage must be between 1 and {len(pdf_reader.pages)}")
    all_sources: set[int] = set()
    style_prefixes = manifest.get("paragraphStylePrefixes")
    if style_prefixes is None:
        style_prefixes = []
    if not isinstance(style_prefixes, list) or not all(isinstance(x, str) for x in style_prefixes):
        _fail("paragraphStylePrefixes must be a string list")
    for spec in normalized:
        supplied = _source_pages(spec, manifest)
        if supplied[-1] > pdf_last_allowed:
            _fail(f"source page range for {spec['id']} exceeds maxSourcePage {pdf_last_allowed}")
        overlap = all_sources.intersection(supplied)
        if overlap:
            _fail(f"source page ranges overlap at {sorted(overlap)}")
        all_sources.update(supplied)
    ids = {a["id"] for a in normalized}
    _validate_existing_catalog(website_root, year, ids)
    _source(root, manifest.get("coverSource", ""), "cover image")
    chapter_assets = manifest.get("chapterAssets", {})
    if not isinstance(chapter_assets, dict):
        _fail("chapterAssets must be an object")
    for chapter, rel in chapter_assets.items():
        _source(root, rel, f"{chapter} chapter image")
    return normalized, workbook, pdf_reader


def run_import(manifest_path: Path, project_root: Path | None = None, website_root: Path | None = None) -> dict[str, Any]:
    manifest_path = manifest_path.resolve()
    root = (project_root or manifest_path.parent.parent).resolve()
    website = (website_root or manifest_path.parent.parent / "website").resolve()
    manifest = _read_json(manifest_path)
    specs, metadata, reader = _validate_manifest(manifest, root, website)
    prefixes = manifest.get("paragraphStylePrefixes", [])
    year = manifest["year"]
    edition: dict[str, Any] = {"schemaVersion": 1, "year": year, "title": manifest["title"], "articles": []}
    staged: dict[Path, bytes] = {}
    audit_summary = []
    source_differences = []
    read_dir = Path("read") / str(year)
    for spec in specs:
        info = metadata[spec["metadataRow"]] if manifest.get("metadataWorkbook") else _inline_metadata(spec)
        pages = _source_pages(spec, manifest)
        supplementary = []
        paragraph_sources = []
        difference_count = 0
        if spec.get("idmlSource"):
            paragraphs = _extract_idml(_source(root, spec["idmlSource"], f"IDML source for {spec['id']}"), prefixes)
            _, audit = reconcile_reading_order(paragraphs, reader, pages)
            matches = sorted((r for r in audit if r["status"] == "matched"), key=lambda r: (r["sourcePage"], r["column"], -r["y"], r["inputIndex"]))
            supplementary = [paragraphs[r["inputIndex"]] for r in audit if r["status"] == "ambiguous"]
            differences = [paragraphs[r["inputIndex"]] for r in audit if r["status"] == "unmatched"]
            difference_count = len(differences)
            if differences:
                source_differences.append({"id": spec["id"], "unpublishedSourceDifferences": differences})
            audit_summary.append({"id": spec["id"], "inputParagraphs": len(paragraphs), "matched": len(matches), "supplementary": len(supplementary), "unpublishedDifferences": difference_count})
            paragraph_sources = [{"sourcePage": r["sourcePage"], "column": r["column"]} for r in matches]
            paragraphs = [paragraphs[r["inputIndex"]] for r in matches]
            if not paragraphs:
                _fail(f"no source-verified body paragraphs for {spec['id']}")
            body_source = "依原刊文字核對並整理跨欄讀序；圖表與原始版面請對照 PDF"
        elif spec.get("paragraphs"):
            paragraphs = [re.sub(r"\s+", " ", p).strip() for p in spec["paragraphs"] if p.strip()]
            body_source = "manifest 提供的完整段落"
        else:
            paragraphs = _pdf_paragraphs(reader, pages)
            body_source = "原始 PDF 文字擷取；請以來源 PDF 校對跨欄閱讀順序"
        source_rel = f"sources/{spec['id']}.pdf"
        article = {
            "id": spec["id"], "year": year, "order": spec["order"], "chapter": spec["chapter"],
            "title": info["title"], "authors": info["authors"], "summary": info["summary"],
            "highlights": info["highlights"], "tags": [],
            "minutes": max(1, math.ceil(sum(len(p) for p in paragraphs) / 500)),
            "printedStart": spec["printedStart"], "printedEnd": spec["printedEnd"],
            "sourcePdf": source_rel, "sourcePages": pages, "paragraphs": paragraphs,
            "supplementaryParagraphs": supplementary, "paragraphSources": paragraph_sources,
            "unpublishedDifferenceCount": difference_count,
            "htmlUrl": (read_dir / f"{spec['id']}.html").as_posix(),
        }
        edition["articles"].append(article)
        staged[website / "public" / read_dir / f"{spec['id']}.html"] = _article_html(article, edition["title"], body_source).encode("utf-8")
        writer = PdfWriter()
        for page_no in pages:
            writer.add_page(reader.pages[page_no - 1])
        output = __import__("io").BytesIO()
        writer.write(output)
        staged[website / "public" / source_rel] = output.getvalue()
    staged[website / "public" / read_dir / "index.html"] = _index_html(edition).encode("utf-8")
    staged[website / "public/data" / f"{year}-source-audit.json"] = (json.dumps(audit_summary, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    if source_differences:
        staged[website / "content" / f"{year}-review.json"] = (json.dumps(source_differences, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    catalog, others = _validate_existing_catalog(website, year, {a["id"] for a in edition["articles"]})
    entry = {"year": year, "title": manifest["title"], "articleCount": len(edition["articles"]), "dataUrl": f"data/{year}.json", "coverUrl": f"assets/cover-{year}.webp"}
    original_years = catalog["years"]
    years: list[dict[str, Any]] = []
    replaced = False
    for old in original_years:
        if old.get("year") == year:
            if not replaced:
                years.append(entry)
                replaced = True
        else:
            years.append(old)
    if not replaced:
        years.append(entry)
    catalog_out = {"schemaVersion": 1, "years": years}
    staged[website / "public/data" / f"{year}.json"] = (json.dumps(edition, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    staged[website / "public/data/catalog.json"] = (json.dumps(catalog_out, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    asset_sources = {f"assets/cover-{year}.webp": manifest["coverSource"]}
    for chapter, source in manifest.get("chapterAssets", {}).items():
        asset_sources[f"assets/chapter-{chapter}-{year}.webp"] = source
    for dest_rel, source_rel in asset_sources.items():
        source = _source(root, source_rel, "image source")
        try:
            with Image.open(source) as im:
                im = ImageOps.exif_transpose(im).convert("RGBA")
                buffer = __import__("io").BytesIO()
                im.save(buffer, format="WEBP", quality=88, method=6)
                staged[website / "public" / dest_rel] = buffer.getvalue()
        except Exception as exc:
            _fail(f"cannot convert image {source}: {exc}")
    for dest, data in staged.items():
        dest.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=dest.parent, prefix=f".{dest.name}.", delete=False) as tmp:
            tmp.write(data)
            temp_path = Path(tmp.name)
        temp_path.replace(dest)
    return {"year": year, "articles": len(edition["articles"]), "paragraphs": sum(len(a["paragraphs"]) for a in edition["articles"]), "outputs": len(staged)}


def main(argv: list[str] | None = None) -> int:
    default_manifest = Path(__file__).resolve().parents[1] / "content/2026.json"
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", type=Path, default=default_manifest, help="edition manifest JSON")
    parser.add_argument("--project-root", type=Path, help="root containing source XLSX/PDF/IDML paths")
    parser.add_argument("--website-root", type=Path, help="website root containing public/")
    args = parser.parse_args(argv)
    try:
        result = run_import(args.manifest, args.project_root, args.website_root)
    except ImportErrorDetail as exc:
        print(f"import error: {exc}", file=sys.stderr)
        return 2
    print(f"Imported {result['articles']} articles for {result['year']}: {result['paragraphs']} paragraphs, {result['outputs']} generated files.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
